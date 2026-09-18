import os
import csv
import json
import requests
import openpyxl
from docx import Document
from dotenv import load_dotenv
from anthropic import Anthropic

# Carga variables de entorno
load_dotenv()

API_URL = os.getenv("API_URL", "http://localhost:8000").rstrip("/")
ADMIN_KEY = os.getenv("ADMIN_API_KEY")
SLUG_EVENTO = os.getenv("SLUG_EVENTO", "cumple-octubre")
ANTHROPIC_KEY = os.getenv("ANTHROPIC_API_KEY")

if not ADMIN_KEY:
    raise ValueError("Falta ADMIN_API_KEY en el archivo .env")

# ----------------- UTILIDADES -----------------

def limpiar_registro(nombre, apellido):
    """Sanitiza strings y limita a 100 caracteres para proteger la BD."""
    return {
        "nombre": str(nombre).strip()[:100],
        "apellido": str(apellido).strip()[:100]
    }

def enviar_en_bloques(invitados, slug, tamano_bloque=50):
    """Envía los invitados en lotes para no saturar el servidor."""
    if not invitados:
        print("No se encontraron invitados para enviar.")
        return

    url = f"{API_URL}/admin/eventos/{slug}/invitados"
    headers = {
        "Content-Type": "application/json",
        "X-Admin-Key": ADMIN_KEY
    }

    print(f"Subiendo {len(invitados)} invitados en bloques de {tamano_bloque}...")
    
    for i in range(0, len(invitados), tamano_bloque):
        bloque = invitados[i : i + tamano_bloque]
        payload = {"invitados": bloque}
        
        try:
            resp = requests.post(url, headers=headers, json=payload)
            if resp.status_code == 200:
                print(f"✅ Lote {i//tamano_bloque + 1} exitoso: {resp.json().get('mensaje')}")
            else:
                print(f"❌ Error en lote {i//tamano_bloque + 1} ({resp.status_code}): {resp.text}")
        except requests.exceptions.RequestException as e:
            print(f"❌ Error de conexión en lote {i//tamano_bloque + 1}:", e)


# ----------------- PROCESAMIENTO ESTRUCTURADO -----------------

def extraer_csv(ruta):
    invitados = []
    with open(ruta, mode='r', encoding='utf-8') as f:
        reader = csv.DictReader(f)
        for row in reader:
            if 'nombre' in row and 'apellido' in row:
                invitados.append(limpiar_registro(row["nombre"], row["apellido"]))
    return invitados

def extraer_excel(ruta):
    invitados = []
    wb = openpyxl.load_workbook(ruta, data_only=True)
    hoja = wb.active
    
    encabezados = {str(cell.value).lower().strip(): i for i, cell in enumerate(hoja[1]) if cell.value}
    
    if 'nombre' not in encabezados or 'apellido' not in encabezados:
        raise ValueError("El Excel debe tener columnas llamadas 'nombre' y 'apellido' en la primera fila")
        
    for row in hoja.iter_rows(min_row=2, values_only=True):
        nombre = row[encabezados['nombre']]
        apellido = row[encabezados['apellido']]
        if nombre and apellido:
            invitados.append(limpiar_registro(nombre, apellido))
    return invitados


# ----------------- PROCESAMIENTO CON IA (NO ESTRUCTURADO) -----------------

def procesar_con_ia(texto):
    """Usa Claude para extraer nombres de texto libre y devolver JSON estricto."""
    if not ANTHROPIC_KEY:
        raise ValueError("Se requiere ANTHROPIC_API_KEY para procesar archivos no estructurados.")
        
    print("Enviando texto a Claude para extracción de entidades...")
    cliente = Anthropic(api_key=ANTHROPIC_KEY)
    
    prompt = f"""
    Extrae todos los nombres de personas mencionados en este texto.
    Devuelve ÚNICAMENTE un arreglo JSON válido donde cada objeto tenga las claves "nombre" y "apellido".
    No incluyas texto adicional ni explicaciones, solo el JSON.
    Si alguien tiene varios nombres/apellidos, agrúpalos lógicamente en los dos campos.
    
    Texto:
    {texto}
    """
    
    respuesta = cliente.messages.create(
        model="claude-3-haiku-20240307",
        max_tokens=1024,
        messages=[{"role": "user", "content": prompt}]
    )
    
    contenido = respuesta.content[0].text
    
    try:
        # Busca el inicio y fin del array JSON por si el LLM añade texto extra
        inicio = contenido.find('[')
        fin = contenido.rfind(']') + 1
        json_limpio = contenido[inicio:fin]
        
        datos = json.loads(json_limpio)
        return [limpiar_registro(d["nombre"], d["apellido"]) for d in datos]
    except (json.JSONDecodeError, ValueError) as e:
        print("Error parseando la respuesta de la IA:", e)
        print("Respuesta cruda:", contenido)
        return []

def leer_docx(ruta):
    doc = Document(ruta)
    texto = "\n".join([p.text for p in doc.paragraphs if p.text.strip()])
    return procesar_con_ia(texto)

def leer_txt(ruta):
    with open(ruta, 'r', encoding='utf-8') as f:
        texto = f.read()
    return procesar_con_ia(texto)


# ----------------- ENRUTADOR PRINCIPAL -----------------

def parsear_archivo(ruta):
    if not os.path.exists(ruta):
        raise FileNotFoundError(f"No se encontró el archivo: {ruta}")
        
    ext = os.path.splitext(ruta)[1].lower()
    
    if ext == '.csv':
        return extraer_csv(ruta)
    elif ext in ['.xlsx', '.xls']:
        return extraer_excel(ruta)
    elif ext == '.docx':
        return leer_docx(ruta)
    elif ext == '.txt':
        return leer_txt(ruta)
    else:
        raise ValueError(f"Formato no soportado aún: {ext}. Soporta: csv, xlsx, docx, txt.")

if __name__ == "__main__":
    # Cambia este nombre por el archivo real que quieras probar
    archivo_prueba = "lista.csv" 
    
    try:
        lista_final = parsear_archivo(archivo_prueba)
        enviar_en_bloques(lista_final, SLUG_EVENTO)
    except Exception as e:
        print(f"Se detuvo la ejecución: {e}")